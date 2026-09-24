from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.tests.test_samples import sample_payload


class MemoryStorage:
    def __init__(self):
        self.objects = {}

    def save(self, content: bytes, extension: str, prefix: str = "originals") -> str:
        key = f"{prefix}/report{extension}"
        self.objects[key] = content
        return key

    def read(self, key: str) -> bytes:
        return self.objects[key]


@pytest.mark.parametrize("value", ["=1+1", "+cmd", "-2+3", "@SUM(A1)"])
def test_excel_user_text_is_not_a_formula(value):
    from backend.app.services.reports import sanitize_excel_text

    assert sanitize_excel_text(value) == "'" + value


def test_report_contains_real_sample_and_is_protected(client, auth_headers, uploaded_image):
    sampled_at = datetime.now().replace(microsecond=0) - timedelta(hours=49)
    payload = sample_payload(uploaded_image["id"], sampled_at, "茶叶蛋")
    sample = client.post(
        "/api/v1/samples",
        json=payload,
        headers={**auth_headers, "Idempotency-Key": "report-sample"},
    ).json()
    client.post(
        f"/api/v1/samples/{sample['id']}/dispose",
        json={"method": "discard", "note": "按规定废弃", "reviewer_name": "张海涛"},
        headers=auth_headers,
    )

    response = client.post(
        "/api/v1/reports",
        json={"date_from": sampled_at.date().isoformat(), "date_to": datetime.now().date().isoformat()},
        headers=auth_headers,
    )

    assert response.status_code == 201
    anonymous = client.get(response.json()["download_url"])
    assert anonymous.status_code == 401
    download = client.get(response.json()["download_url"], headers=auth_headers)
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content))
    sheet = workbook.active
    assert sheet["A2"].value == "茶叶蛋"
    assert sheet["K2"].value == "张海涛"


def test_report_rejects_reversed_date_range(client, auth_headers):
    response = client.post(
        "/api/v1/reports",
        json={"date_from": "2026-09-10", "date_to": "2026-09-01"},
        headers=auth_headers,
    )

    assert response.status_code == 422
    assert response.json()["code"] == "INVALID_DATE_RANGE"


def test_report_download_uses_storage_adapter(client, auth_headers, monkeypatch):
    storage = MemoryStorage()
    monkeypatch.setattr("backend.app.routes.reports.storage_for", lambda _settings: storage)

    response = client.post(
        "/api/v1/reports",
        json={"date_from": "2026-09-10", "date_to": "2026-09-10"},
        headers=auth_headers,
    )

    assert response.status_code == 201
    report_key = next(iter(storage.objects))
    assert report_key.startswith("reports/")

    download = client.get(response.json()["download_url"], headers=auth_headers)
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content))
    assert workbook.active.title == "留样台账"

from datetime import datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from ..models import AuditEvent, RecognitionRun, Report, Sample, User
from .samples import append_audit, status_for
from .dashboard import current_local_time


def sanitize_excel_text(value: str) -> str:
    return "'" + value if value and value[0] in "=+-@" else value


def build_report_workbook(db: Session, samples: list[Sample], timezone_name: str) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "留样台账"
    headers = [
        "菜名", "餐次", "留样时间", "到期时间", "留样人", "留样量(g)",
        "温度(℃)", "状态", "处置方式", "处置说明", "复核人", "图片哈希", "审计编号",
    ]
    sheet.append(headers)
    now = current_local_time(timezone_name)
    for sample in samples:
        keeper = db.get(User, sample.keeper_id)
        image_hash = sample.images[0].sha256 if sample.images else ""
        audit_ids = [str(item.id) for item in db.query(AuditEvent).filter(
            AuditEvent.entity_type == "sample", AuditEvent.entity_id == sample.id
        ).all()]
        disposal = sample.disposal
        sheet.append(
            [
                sanitize_excel_text(sample.dish_name),
                sample.meal,
                sample.sampled_at,
                sample.expires_at,
                sanitize_excel_text(keeper.display_name),
                sample.amount_g,
                sample.temperature_c,
                status_for(sample, now),
                disposal.method if disposal else "",
                sanitize_excel_text(disposal.note) if disposal else "",
                sanitize_excel_text(disposal.reviewer_name) if disposal else "",
                image_hash,
                ",".join(audit_ids),
            ]
        )
    fill = PatternFill("solid", fgColor="0F3E17")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    widths = [24, 12, 20, 20, 14, 12, 12, 14, 16, 28, 14, 68, 18]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    return workbook


def create_report(db: Session, actor: User, date_from, date_to, storage, timezone_name: str) -> Report:
    start = datetime.combine(date_from, time.min)
    end = datetime.combine(date_to + timedelta(days=1), time.min)
    samples = (
        db.query(Sample)
        .filter(Sample.sampled_at >= start, Sample.sampled_at < end)
        .order_by(Sample.sampled_at.asc(), Sample.id.asc())
        .all()
    )
    workbook = build_report_workbook(db, samples, timezone_name)
    output = BytesIO()
    workbook.save(output)
    report = Report(
        owner_id=actor.id,
        date_from=date_from,
        date_to=date_to,
        storage_key=storage.save(output.getvalue(), ".xlsx", prefix="reports"),
    )
    db.add(report)
    db.flush()
    append_audit(
        db,
        "report",
        report.id,
        "exported",
        actor.id,
        None,
        {"date_from": date_from.isoformat(), "date_to": date_to.isoformat(), "rows": len(samples)},
    )
    db.commit()
    db.refresh(report)
    return report
